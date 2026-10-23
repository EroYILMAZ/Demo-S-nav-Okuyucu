"""Okunan puanları Excel dosyasına dönüştürür."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import sorular

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _sort_key(rec: dict):
    no = rec.get("ogrenci_no") or ""
    return (0, int(no), "") if no.isdigit() else (1, 0, rec.get("ad_soyad") or "")


def build_workbook(exam: dict) -> bytes:
    """exam: matching.annotate() ile uyarıları eklenmiş sınav."""
    maxes = exam["tam_puanlar"]
    n = len(maxes)
    records = [r for r in exam["kayitlar"] if r.get("durum") == "tamam"]
    # okunamamış / sırada bekleyen kağıtlar da durumlarıyla yer alsın (öğrenci "kağıdı yok" sanılmasın)
    pending = [r for r in exam["kayitlar"] if r.get("durum") != "tamam"]
    # yeniden okunan ya da okunamayan kağıdın öğrencisi "kağıdı okutulmadı" görünmesin (kağıdı var, puanı bekliyor)
    read_nos = {r.get("ogrenci_no") for r in exam["kayitlar"] if r.get("ogrenci_no")}
    for s in exam.get("ogrenci_listesi", []):
        if s["no"] not in read_nos:
            records.append({"ogrenci_no": s["no"], "ad_soyad": s["ad"], "puanlar": [], "eksik": True})
    records.sort(key=_sort_key)
    for r in pending:
        state = {"hata": f"Okunamadı: {r.get('hata') or ''}".strip(), "eksik": "Diğer sayfası bekleniyor"}.get(
            r.get("durum"), "Henüz okunmadı (sırada)")
        records.append({"ogrenci_no": r.get("ogrenci_no") or "", "ad_soyad": r.get("ad_soyad") or "", "puanlar": [],
                        "eksik": True, "not": state})

    wb = Workbook()
    ws = wb.active
    ws.title = "Puanlar"

    title = " - ".join(x for x in (exam.get("sinif"), exam.get("ders"), exam.get("ad")) if x)
    ws["A1"] = title or "Sınav Puanları"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Oluşturma: {datetime.now():%d.%m.%Y %H:%M}   ·   Sarı hücreler kontrol edilmeli (açıklama için hücre notuna bakın)"
    ws["A2"].font = Font(italic=True, color="7F7F7F", size=9)

    hr = 4
    headers = ["Sıra", "Okul No", "Adı Soyadı"] + [f"S{lab} ({m:g})" for lab, m in zip(sorular.labels_of(exam), maxes)] + [
        f"Toplam ({sum(maxes):g})",
        "Kontrol",
    ]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    first_q_col = 4
    last_q_col = first_q_col + n - 1
    total_col = last_q_col + 1
    note_col = total_col + 1

    for i, rec in enumerate(records):
        r = hr + 1 + i
        cells_warn = rec.get("hucre_uyarilari", {})
        no = rec.get("ogrenci_no") or ""
        # baştaki sıfırlar kaybolmasın: "0368" gibi numaralar metin olarak kalır
        values = [i + 1, int(no) if no.isdigit() and not no.startswith("0") else no, rec.get("ad_soyad") or ""]
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=v)
        confirmed = bool(rec.get("onaylandi"))
        if confirmed:
            cells_warn = {}  # öğretmen onayladı: sarı işaret ve not gerekmez
        for key, col in (("no", 2), ("ad", 3)):
            if key in cells_warn:
                _warn(ws.cell(row=r, column=col), cells_warn[key])

        puanlar = rec.get("puanlar") or []
        for q in range(n):
            p = puanlar[q] if q < len(puanlar) else None
            cell = ws.cell(row=r, column=first_q_col + q, value=_num(p))
            if f"p{q}" in cells_warn:
                _warn(cell, cells_warn[f"p{q}"])

        if rec.get("eksik"):
            ws.cell(row=r, column=note_col, value=rec.get("not") or "Kağıdı okutulmadı")
            for c in range(1, note_col + 1):
                ws.cell(row=r, column=c).border = BORDER
                ws.cell(row=r, column=c).font = Font(color="7F7F7F")
            continue

        ws.cell(
            row=r,
            column=total_col,
            value=f"=SUM({get_column_letter(first_q_col)}{r}:{get_column_letter(last_q_col)}{r})",
        ).font = Font(bold=True)
        notes = [] if confirmed else list(rec.get("uyarilar", [])) + list(cells_warn.values())
        ws.cell(row=r, column=note_col, value="; ".join(notes) if notes else ("Onaylandı" if confirmed else ""))
        if rec.get("uyarilar") and not confirmed:
            _warn(ws.cell(row=r, column=total_col), "\n".join(rec["uyarilar"]))

        for c in range(1, note_col + 1):
            ws.cell(row=r, column=c).border = BORDER
            if c != 3 and c != note_col:
                ws.cell(row=r, column=c).alignment = Alignment(horizontal="center")

    if records:
        r = hr + len(records) + 1
        ws.cell(row=r, column=3, value="Sınıf Ortalaması").font = Font(bold=True)
        for col in range(first_q_col, total_col + 1):
            L = get_column_letter(col)
            cell = ws.cell(row=r, column=col, value=f"=IFERROR(ROUND(AVERAGE({L}{hr + 1}:{L}{r - 1}),2),\"\")")
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
            cell.border = BORDER

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 26
    for col in range(first_q_col, total_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 9
    ws.column_dimensions[get_column_letter(note_col)].width = 50
    ws.row_dimensions[hr].height = 30
    ws.freeze_panes = ws.cell(row=hr + 1, column=first_q_col)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _num(v):
    if v is None:
        return None
    return int(v) if float(v).is_integer() else v


def _warn(cell, text: str) -> None:
    cell.fill = WARN_FILL
    cell.comment = Comment(text, "Sınav Okuyucu")
