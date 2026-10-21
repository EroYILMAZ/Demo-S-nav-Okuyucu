"""Sınıf listesiyle eşleştirme ve okunan puanların kontrolü."""
import csv
import io
import re
from difflib import SequenceMatcher

_TR_MAP = str.maketrans("çğıöşüâîû", "cgiosuaiu")


def tr_norm(text: str) -> str:
    """Türkçe harf farklarını ve büyük/küçük harfi yok sayan karşılaştırma anahtarı."""
    text = (text or "").replace("İ", "i").replace("I", "ı").lower().translate(_TR_MAP)
    return " ".join(re.sub(r"[^a-z0-9 ]", " ", text).split())


def tr_title(text: str) -> str:
    """'ALİ YILMAZ' -> 'Ali Yılmaz' (Türkçe i/ı kurallarıyla)."""
    def word(w: str) -> str:
        if not w:
            return w
        first = w[0].replace("i", "İ").replace("ı", "I").upper()
        rest = w[1:].replace("I", "ı").replace("İ", "i").lower()
        return first + rest
    return " ".join(word(w) for w in (text or "").split())


def name_similarity(a: str, b: str) -> float:
    a, b = tr_norm(a), tr_norm(b)
    if not a or not b:
        return 0.0
    direct = SequenceMatcher(None, a, b).ratio()
    # sıra farkını (Soyad Ad) tolere et
    sorted_ = SequenceMatcher(None, " ".join(sorted(a.split())), " ".join(sorted(b.split()))).ratio()
    return max(direct, sorted_)


def parse_roster_text(text: str) -> list[dict]:
    """Yapıştırılan listeyi okur. Her satır: '123 Ali Yılmaz', '123;Ali Yılmaz' veya '123<TAB>Ali<TAB>Yılmaz'."""
    roster = []
    for line in (text or "").splitlines():
        tokens = [t for t in re.split(r"[\s;,|]+", line.strip()) if t]
        # baştaki sayılar: "sıra no" + "okul no" olabilir; öğrenci numarası sondaki sayıdır
        nums = []
        while tokens and re.fullmatch(r"\d+[.)-]?", tokens[0]):
            nums.append(re.sub(r"\D", "", tokens.pop(0)))
        words = [t for t in tokens if re.search(r"[^\W\d_]", t)]
        if nums and words:
            roster.append({"no": nums[-1], "ad": tr_title(" ".join(words))})
    return _dedupe(roster)


def parse_roster_rows(rows: list[list]) -> list[dict]:
    """Excel/CSV satırlarından listeyi okur. Başlık satırında 'no' ve 'ad' sütunları aranır."""
    rows = [[("" if c is None else str(c).strip()) for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    no_col = name_cols = None
    start = 0
    for i, r in enumerate(rows[:15]):
        keys = [tr_norm(c) for c in r]
        nc = next((j for j, k in enumerate(keys) if k in ("no", "okul no", "numara", "ogrenci no", "okul numarasi")), None)
        if nc is None:
            nc = next((j for j, k in enumerate(keys)
                       if (k.endswith(" no") or k.startswith("no ")) and not k.startswith(("sira", "s no"))), None)
        full = next((j for j, k in enumerate(keys) if k in ("adi soyadi", "ad soyad", "ad soyadi", "adi ve soyadi", "ogrenci adi soyadi", "isim")), None)
        first = next((j for j, k in enumerate(keys) if k in ("adi", "ad")), None)
        last = next((j for j, k in enumerate(keys) if k in ("soyadi", "soyad")), None)
        if nc is not None and (full is not None or first is not None):
            no_col = nc
            name_cols = [full] if full is not None else [c for c in (first, last) if c is not None]
            start = i + 1
            break
    if no_col is None:
        # başlık yoksa: 'no' ilk sayısal sütun, isim onu izleyen metin sütunları
        return parse_roster_text("\n".join("\t".join(r) for r in rows))

    roster = []
    for r in rows[start:]:
        no = re.sub(r"\D", "", r[no_col] if no_col < len(r) else "")
        name = " ".join(r[c] for c in name_cols if c < len(r) and r[c])
        if no and name:
            roster.append({"no": no, "ad": tr_title(name)})
    return _dedupe(roster)


def _decode(content: bytes) -> str:
    """UTF-8 değilse Türkçe Windows kodlaması (Excel'in Türkçe CSV'si cp1254 kaydeder)."""
    for enc in ("utf-8-sig", "cp1254"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def parse_roster_file(filename: str, content: bytes) -> list[dict]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            return parse_roster_rows([list(r) for r in ws.iter_rows(values_only=True)])
        except Exception:
            raise ValueError("Excel dosyası açılamadı. Dosyayı Excel'de açıp yeniden kaydedin ya da listeyi kopyalayıp yapıştırın.")
    if name.endswith(".xls"):
        raise ValueError("Eski .xls biçimi okunamıyor. Dosyayı Excel'de açıp .xlsx olarak kaydedin ya da listeyi kopyalayıp yapıştırın.")
    text = _decode(content)
    if name.endswith(".csv"):
        try:
            dialect = csv.Sniffer().sniff(text[:2000], delimiters=";,\t") if text.strip() else csv.excel
            return parse_roster_rows(list(csv.reader(io.StringIO(text), dialect)))
        except csv.Error:
            return parse_roster_text(text)  # ayırıcı anlaşılamadı: satır satır oku
    return parse_roster_text(text)


def _dedupe(roster: list[dict]) -> list[dict]:
    seen, out = set(), []
    for s in roster:
        if s["no"] not in seen:
            seen.add(s["no"])
            out.append(s)
    return out


def _one_digit_apart(a: str, b: str) -> bool:
    return bool(a) and len(a) == len(b) and sum(x != y for x, y in zip(a, b)) <= 1


def resolve_student(no: str, name: str, roster: list[dict], p_no: float = 0.9, choice: dict | None = None):
    """Okunan numara/ismi sınıf listesine eşler ve öğrenci kimliğinin eminliğini (0-1) hesaplar.

    choice: modelin başlığı sınıf listesiyle karşılaştırıp seçtiği öğrenci {"no", "eminlik"} (varsa).
    Dönüş: (no, ad, uyarı_notu, bilgi_notu, eminlik). Uyarı notu kağıdı kontrole düşürür; bilgi notu yalnızca açıklar.
    """
    if not roster:
        return no, tr_title(name), "", "", (p_no if no else 0.0)

    by_no = next((s for s in roster if s["no"] == no), None) if no else None
    best = max(roster, key=lambda s: name_similarity(name, s["ad"]))
    best_sim = name_similarity(name, best["ad"])
    ch_no = (choice or {}).get("no")
    ch_conf = float((choice or {}).get("eminlik") or 0.0)
    chosen = next((s for s in roster if s["no"] == ch_no), None) if ch_no else None

    if by_no:
        sim = name_similarity(name, by_no["ad"])
        if sim >= 0.6:
            conf = 1 - (1 - p_no) * (1 - min(0.95, sim))
            if chosen and chosen["no"] != by_no["no"] and ch_conf >= 0.8:
                return by_no["no"], by_no["ad"], (
                    f"Numara ve isim {by_no['ad']} diyor ama listeyle karşılaştırma {chosen['ad']} gösteriyor."), "", min(conf, 0.6)
            return by_no["no"], by_no["ad"], "", "", conf
        if chosen and chosen["no"] == by_no["no"]:
            conf = 1 - (1 - p_no) * (1 - ch_conf)
            return by_no["no"], by_no["ad"], "", (
                f"İsim “{name}” diye okundu; numara ve sınıf listesi karşılaştırması aynı öğrenciyi gösteriyor."), conf
        if chosen and chosen["no"] != by_no["no"]:
            return by_no["no"], by_no["ad"], (
                f"Numara {no} ({by_no['ad']}) ama yazılan isim listedeki {chosen['ad']} gibi görünüyor."), "", 0.5
        if best_sim >= 0.8 and best["no"] != by_no["no"]:
            return best["no"], best["ad"], (
                f"Numara {no} ({by_no['ad']}) ile isim “{name}” farklı öğrencileri gösteriyor; isme göre seçildi."), "", 0.5
        return by_no["no"], by_no["ad"], "", f"İsim “{name}” diye okundu, numaraya göre eşleştirildi.", p_no * 0.9

    reason = "okunamadı" if not no else f"{no} listede yok"
    if chosen:
        sim = name_similarity(name, chosen["ad"])
        conf = ch_conf * (0.85 + 0.15 * min(1.0, sim / 0.7))
        if _one_digit_apart(no, chosen["no"]):
            conf = 1 - (1 - conf) * 0.5  # okunan numara seçilen öğrencinin numarasına çok yakın: ek kanıt
        # ikinci aday: okunan el yazısı isme en çok benzeyen (seçilen dışındaki) öğrenci; isim yoksa modelin rakam olasılığı
        others = sorted((s for s in roster if s["no"] != chosen["no"]), key=lambda s: -name_similarity(name, s["ad"]))
        other = others[0] if others and name and name_similarity(name, others[0]["ad"]) >= 0.5 else \
            next((s for s in roster if s["no"] == (choice or {}).get("alternatif")), None)
        maybe = f" Diğer olasılık: {other['no']} {other['ad']}." if other and conf < 0.9 else ""
        return chosen["no"], chosen["ad"], "", f"Numara {reason}; öğrenci, sınıf listesiyle karşılaştırılarak bulundu.{maybe}", conf
    if best_sim >= 0.7:
        conf = 0.5 + 0.4 * (best_sim - 0.7) / 0.3
        return best["no"], best["ad"], f"Numara {reason}; isme göre eşleştirildi.", "", conf
    return no, tr_title(name), "Öğrenci sınıf listesinde bulunamadı.", "", 0.2


def match_student(no: str, name: str, roster: list[dict]) -> tuple[str, str, str]:
    """Geriye dönük uyumluluk: (no, ad, uyarı_notu)."""
    r = resolve_student(no, name, roster)
    return r[0], r[1], r[2]


def record_warnings(rec: dict, exam: dict, duplicates: set[str]) -> tuple[list[str], dict[str, str]]:
    """Bir kaydın satır uyarıları ve hücre uyarıları ({'p0': '...', 'no': '...'}).
    Kontrol, gerekçeli uyarılara dayanır (okumalar farklı, puan bulunamadı, toplam tutmuyor, öğrenci belirsiz…);
    yüzdelik bir "eminlik" puanı kullanılmaz."""
    row: list[str] = []
    cells: dict[str, str] = {}
    if rec.get("durum") != "tamam":
        return row, cells

    maxes = exam["tam_puanlar"]
    puanlar = rec.get("puanlar") or []
    supheli = rec.get("supheli") or {}
    for i, mx in enumerate(maxes):
        p = puanlar[i] if i < len(puanlar) else None
        if str(i) in supheli:
            cells[f"p{i}"] = supheli[str(i)]
        if p is None:
            # şablon tanılaması ("kutu yok", "sayfa bulunamadı", "puan bulunamadı") daha açıklayıcıdır; ezme
            cells.setdefault(f"p{i}", "Puan okunamadı")
        elif p < 0:
            cells[f"p{i}"] = "Negatif puan"
        elif p > mx:
            cells[f"p{i}"] = f"Tam puandan ({mx:g}) büyük"


    total = sum(p for p in puanlar if p is not None)
    yazili = rec.get("yazili_toplam")
    if yazili is not None and abs(yazili - total) > 0.01 and not rec.get("toplam_uyarisi_kapali"):
        row.append(f"Kağıttaki toplam ({yazili:g}) ile soru puanlarının toplamı ({total:g}) tutmuyor.")

    no = rec.get("ogrenci_no") or ""
    if not no:
        cells["no"] = "Numara okunamadı"
    elif no in duplicates:
        cells["no"] = "Bu numara birden fazla kağıtta var"
    elif exam["ogrenci_listesi"] and not any(s["no"] == no for s in exam["ogrenci_listesi"]):
        cells["no"] = "Sınıf listesinde yok"
    k = rec.get("eminlik_kimlik")  # iç ölçü: numara net okunamadıysa ya da öğrenci listeden tahmin edildiyse düşük
    if k is not None and k < 0.9 and "no" not in cells:
        cells["no"] = "Öğrenci kesin belirlenemedi, kontrol edin"
    if not rec.get("ad_soyad"):
        cells["ad"] = "İsim okunamadı"
    if rec.get("eslesme_notu"):
        row.append(rec["eslesme_notu"])
    if rec.get("hizalanamayan"):
        sayfalar = ", ".join(map(str, rec["hizalanamayan"]))
        row.append(f"{sayfalar}. fotoğraf şablona oturtulamadı (başka baskı/grup ya da bulanık fotoğraf olabilir).")
    per = int(exam.get("sayfa_sayisi", 1))
    if rec.get("eksik_sayfa") and per > 1:
        row.append(f"Sayfası eksik okundu (bu sınav {per} sayfa). Eksik sayfayı ekleyip yeniden okutabilirsiniz.")
    return row, cells


def annotate(exam: dict) -> dict:
    """Sınavın kayıtlarına uyarıları ekleyip döndürür (diske yazılmaz)."""
    counts: dict[str, int] = {}
    for r in exam["kayitlar"]:
        if r.get("durum") == "tamam" and r.get("ogrenci_no"):
            counts[r["ogrenci_no"]] = counts.get(r["ogrenci_no"], 0) + 1
    duplicates = {n for n, c in counts.items() if c > 1}
    for r in exam["kayitlar"]:
        r["uyarilar"], r["hucre_uyarilari"] = record_warnings(r, exam, duplicates)
    return exam
